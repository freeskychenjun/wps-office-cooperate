#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Assertions for `ppt chart` (tests/ppt).

The point of the command is that the numbers are REAL, not a picture: so the checks go
through the saved file with python-pptx, which reads the chart's own cached data.

  python check_chart.py --dir <out-chart>
"""
import argparse
import json
import os
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

_failures = []


def fail(message):
    _failures.append(message)


def load(folder, step):
    path = os.path.join(folder, "resp-%s.json" % step)
    if not os.path.exists(path):
        fail("missing response %r" % step)
        return None, {}
    with open(path, "r", encoding="utf-8") as fh:
        raw = fh.read()
    try:
        doc = json.loads(raw)
    except Exception as exc:
        fail("response %r is not JSON: %s" % (step, exc))
        return None, {}
    return doc, doc.get("result", {})


def ok(step, doc):
    if doc is None:
        return False
    if not doc.get("ok"):
        fail("step %r failed: %s" % (step, json.dumps(doc.get("error"), ensure_ascii=False)))
        return False
    return True


def check_responses(folder):
    doc, single = load(folder, "chart-single")
    if ok("chart-single", doc):
        if single.get("chartType") != 51:
            fail("column chart reported type %r" % single.get("chartType"))
        if single.get("series") != 1:
            fail("one-series payload produced %r series (WPS seeds 3, the extras must be deleted)"
                 % single.get("series"))
        values = single.get("firstSeriesValues") or []
        if values != [1200.0, 1580.0, 2248.0]:
            fail("first series reads %r, expected the payload values" % values)
        if single.get("categories") != ["2020", "2021", "2022"]:
            fail("categories came back as %r" % single.get("categories"))

    doc, two = load(folder, "chart-two")
    if ok("chart-two", doc):
        if two.get("series") != 2:
            fail("two-series payload produced %r series" % two.get("series"))
        if two.get("chartType") != 57:
            fail("barh chart reported type %r" % two.get("chartType"))

    doc, line = load(folder, "chart-line")
    if ok("chart-line", doc):
        if line.get("chartType") != 4:
            fail("line chart reported type %r" % line.get("chartType"))
        values = line.get("firstSeriesValues") or []
        if len(values) != 5:
            fail("line chart has %d points, expected 5" % len(values))

    doc, bad = load(folder, "chart-bad")
    if doc is None or doc.get("ok"):
        fail("an unknown -Type should be refused")
    else:
        if "column|barh|line" not in json.dumps(doc.get("error"), ensure_ascii=False):
            fail("the unknown-type error does not list the accepted types")

    doc, listed = load(folder, "chart-list")
    if ok("chart-list", doc):
        if listed.get("total") != 3:
            fail("chart list found %r charts, expected 3" % listed.get("total"))


def check_saved_deck(folder):
    plain = os.path.join(folder, "deck.plain.pptx")
    if not os.path.exists(plain):
        fail("deck.plain.pptx missing (decrypt step)")
        return
    from pptx import Presentation
    try:
        prs = Presentation(plain)
    except Exception as exc:
        fail("could not read the saved deck: %s" % exc)
        return

    charts = []
    for index, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if getattr(shape, "has_chart", False) and shape.has_chart:
                charts.append((index, shape.chart))
    if len(charts) != 3:
        fail("the saved deck has %d charts, expected 3" % len(charts))
        return

    expected = [
        (1, "2020", 1200.0, "Measured flow"),
        (2, "一月", 120.0, "Measured vs design"),
        (3, "1", 12.1, "Stage"),
    ]
    for (slide_no, chart), (want_slide, want_cat, want_value, want_title) in zip(charts, expected):
        if slide_no != want_slide:
            fail("chart sits on slide %d, expected %d" % (slide_no, want_slide))
        cats = []
        try:
            cats = [str(c) for c in chart.plots[0].categories]
        except Exception as exc:
            fail("slide %d: could not read categories: %s" % (slide_no, exc))
        if cats and cats[0] != want_cat:
            fail("slide %d categories start with %r, expected %r" % (slide_no, cats[0], want_cat))
        values = []
        try:
            values = [round(float(v), 2) for v in chart.series[0].values if v is not None]
        except Exception as exc:
            fail("slide %d: could not read values: %s" % (slide_no, exc))
        if values and abs(values[0] - want_value) > 0.01:
            fail("slide %d first value is %r, expected %r" % (slide_no, values[0], want_value))
        # the payload's point count must be exact: a seeded 4th row would show up here
        want_points = {1: 3, 2: 3, 3: 5}[want_slide]
        if len(values) != want_points:
            fail("slide %d has %d plotted points, expected %d (leftover seed rows must be blanked)"
                 % (slide_no, len(values), want_points))
        try:
            title = chart.chart_title.text_frame.text
        except Exception:
            title = ""
        if want_title not in (title or ""):
            fail("slide %d chart title is %r, expected %r" % (slide_no, title, want_title))


def main(argv):
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", required=True)
    args = ap.parse_args(argv)
    folder = os.path.abspath(args.dir)
    if not os.path.isdir(folder):
        print("no such directory: %s" % folder)
        return 2

    check_responses(folder)
    check_saved_deck(folder)

    if _failures:
        print("")
        for line in _failures:
            print("  FAIL  %s" % line)
        print("\n%d check(s) failed" % len(_failures))
        return 1
    print("\nall checks passed")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
