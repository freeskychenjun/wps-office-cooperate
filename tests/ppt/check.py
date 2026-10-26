#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Acceptance checks for `ppt replace` (tests/ppt).

Everything here is computed from the *pristine* fixture deck with python-pptx, i.e.
by an independent implementation of "what the deck contains".  The driver
(run_replace.ps1) only starts WPS, calls the skill and drops the JSON responses
next to the deck; this script decides whether the run was correct.

  python check.py --dir <out>            check every scenario folder under <out>
  python check.py --dir <out> --list     print the expectations only

A scenario folder holds
  pristine.pptx        ground truth (never opened by WPS)
  deck.pptx            the deck the skill worked on
  resp-<step>.json     one file per skill call that we care about

Assertions per scenario
  dryrun      read-only: counted hits, `applied` false, no snapshot, saved deck
              still byte-identical in *text* to the pristine one
  full        every hit applied; hit text / where-multiset / slots-touched match;
              per-character formatting (size, latin, east-asian face, colour) is
              unchanged outside the replaced spans and taken from the first
              character of the span inside them; nothing skipped
  idem        a second identical run changes nothing
  restore     the slot snapshot takes the whole deck back to pristine
  notes/master/slidesonly/slide1/regex/case/max/items/footer
              the filter narrows the hit set exactly as computed here
"""
import argparse
import json
import os
import re
import sys

from pptx import Presentation
from pptx.oxml.ns import qn

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

TOKEN = "2024"
NEW = "2025"
# The scope a bare `replace` searches: the slide master is opt-in on purpose, so it
# must NOT be in the default expectation.
DEFAULT_WHERES = ("body", "group", "table", "notes", "footer")
# Compared strictly.  `bold` is reported separately: WPS rewrites b="0" as an
# absent attribute on save, which is a spelling change, not a formatting change.
FMT_STRICT = 4

_failures = []
_infos = []


def fail(scenario, message):
    _failures.append("%s: %s" % (scenario, message))


def info(scenario, message):
    _infos.append("%s: %s" % (scenario, message))


# ---------------------------------------------------------------- scanning

def sig_of(run):
    font = run.font
    size = float(font.size.pt) if font.size is not None else None
    if size is None:
        # DrawingML default when sz is absent: 18 pt.  WPS writes exactly that by
        # omitting the attribute, so "no sz" and "sz=1800" have to compare equal.
        size = 18.0
    rpr = font._rPr
    east_asian = None
    el = rpr.find(qn("a:ea")) if rpr is not None else None
    if el is not None:
        east_asian = el.get("typeface")
    color = None
    try:
        if font.color is not None and font.color.type is not None:
            color = str(font.color.rgb)
    except Exception:
        color = None
    bold = None if font.bold is None else bool(font.bold)
    return (size, font.name, east_asian, color, bold)


def para_chars(para):
    chars = []
    for run in para.runs:
        s = sig_of(run)
        for ch in run.text:
            chars.append((ch, s))
    return chars


def slot_from_text_frame(slide_no, where, key, text_frame):
    paras = []
    for para in text_frame.paragraphs:
        paras.append(para_chars(para))
    if not any(paras):
        # A date / slide-number placeholder holds a field (<a:fld>), so text_frame.text
        # is not empty while there is not a single run to format, and WPS drops the
        # field on save.  Such a placeholder is not a text slot.
        return None
    return {"slide": slide_no, "where": where, "key": key, "paras": paras}


def walk_shapes(shapes, slide_no, where, prefix, out):
    for shape in shapes:
        name = shape.name
        if shape.shape_type is not None and int(shape.shape_type) == 6:      # group
            walk_shapes(shape.shapes, slide_no, "group", prefix + name + "/", out)
            continue
        if getattr(shape, "has_table", False) and shape.has_table:
            table = shape.table
            for r, row in enumerate(table.rows):
                for c, cell in enumerate(row.cells):
                    slot = slot_from_text_frame(
                        slide_no, "table", "%s%s[%d,%d]" % (prefix, name, r + 1, c + 1),
                        cell.text_frame)
                    if slot is not None:
                        out.append(slot)
            continue
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            slot = slot_from_text_frame(slide_no, where, prefix + name, shape.text_frame)
            if slot is not None:
                out.append(slot)


def scan(path):
    prs = Presentation(path)
    out = []
    for index, slide in enumerate(prs.slides, start=1):
        walk_shapes(slide.shapes, index, "body", "", out)
        # Only the notes body placeholder carries prose; other shapes on a notes page
        # are the slide thumbnail and (in some templates) prompt text.
        try:
            notes_tf = slide.notes_slide.notes_text_frame
        except Exception:
            notes_tf = None
        if notes_tf is not None and notes_tf.text.strip() != "":
            slot = slot_from_text_frame(index, "notes", "notes", notes_tf)
            if slot is not None:
                out.append(slot)
    master = prs.slide_master
    walk_shapes(master.shapes, 0, "master", "master:", out)
    # Layouts are deliberately NOT scanned: WPS rewrites the placeholder prompt text
    # of a layout when it saves, which has nothing to do with the replace.  The
    # fixture puts no token in a layout, so nothing is being skipped here.
    return out


def slot_text(slot):
    return "\r".join("".join(ch for ch, _ in para) for para in slot["paras"])


# ---------------------------------------------------------------- expectations

def transform(chars, token, new):
    """Apply the replace the skill is supposed to do: non-overlapping, case
    insensitive, new characters take the signature of the span's first character."""
    out = []
    i = 0
    n = len(chars)
    lowered = token.lower()
    while i < n:
        window = "".join(c for c, _ in chars[i:i + len(token)])
        if len(window) == len(token) and window.lower() == lowered:
            sig = chars[i][1]
            for ch in new:
                out.append((ch, sig))
            i += len(token)
        else:
            out.append(chars[i])
            i += 1
    return out


def expected_deck(slots, wheres=DEFAULT_WHERES, token=TOKEN, new=NEW):
    """The deck we expect after a bare `replace`: slots outside the search scope come
    back untouched, which is also how the master-opt-in rule gets asserted."""
    result = []
    for slot in slots:
        if wheres is not None and slot["where"] not in wheres:
            paras = slot["paras"]
        else:
            paras = [transform(para, token, new) for para in slot["paras"]]
        result.append({"slide": slot["slide"], "where": slot["where"],
                       "key": slot["key"], "paras": paras})
    return result


def count_hits(slots, token=TOKEN, case_sensitive=False, wheres=None):
    total = 0
    by_where = {}
    by_slide = {}
    slots_hit = 0
    for slot in slots:
        if wheres is not None and slot["where"] not in wheres:
            continue
        text = slot_text(slot)
        hay = text if case_sensitive else text.lower()
        needle = token if case_sensitive else token.lower()
        n = hay.count(needle)
        if n:
            total += n
            by_where[slot["where"]] = by_where.get(slot["where"], 0) + n
            by_slide[str(slot["slide"])] = by_slide.get(str(slot["slide"]), 0) + n
            slots_hit += 1
    return {"total": total, "byWhere": by_where, "bySlide": by_slide, "slotsTouched": slots_hit}


def scoped_slots(slots, wheres=DEFAULT_WHERES):
    return [s for s in slots if s["where"] in wheres]


def regex_hits(slots, pattern):
    total = 0
    for slot in slots:
        total += len(re.findall(pattern, slot_text(slot)))
    return total


def compare_decks(scenario, expected_slots, actual_slots, strict_format=True):
    """expected_slots is what the deck should look like now; actual_slots is what the
    saved file contains."""
    if actual_slots is None:
        return          # no decrypted copy to look at; reported separately
    if len(actual_slots) != len(expected_slots):
        fail(scenario, "slot count changed: %d -> %d (shapes were added or lost)"
             % (len(expected_slots), len(actual_slots)))
        return
    for want, got in zip(expected_slots, actual_slots):
        label = "%s#%s#%s" % (got["slide"], got["where"], got["key"])
        # The slide master carries template prompt text that WPS may rewrite on save;
        # only the master slot that actually holds the token is worth asserting.
        if want["slide"] == 0 and TOKEN not in slot_text(want):
            continue
        if got["where"] != want["where"] or got["slide"] != want["slide"]:
            fail(scenario, "%s: moved to %s#%s" % (label, got["slide"], got["where"]))
            continue
        if slot_text(got) != slot_text(want):
            fail(scenario, "%s: text is %r, expected %r"
                 % (label, slot_text(got), slot_text(want)))
            continue
        if not strict_format:
            continue
        if len(got["paras"]) != len(want["paras"]):
            fail(scenario, "%s: paragraph count %d != %d"
                 % (label, len(got["paras"]), len(want["paras"])))
            continue
        for pi, (gpara, wpara) in enumerate(zip(got["paras"], want["paras"])):
            if len(gpara) != len(wpara):
                fail(scenario, "%s p%d: %d characters, expected %d"
                     % (label, pi + 1, len(gpara), len(wpara)))
                continue
            for ci, ((gch, gsig), (ech, esig)) in enumerate(zip(gpara, wpara)):
                if gch != ech:
                    fail(scenario, "%s p%d c%d: character %r != %r"
                         % (label, pi + 1, ci + 1, gch, ech))
                    break
                if gsig[:FMT_STRICT] != esig[:FMT_STRICT]:
                    fail(scenario, "%s p%d c%d %r: formatting %r, expected %r"
                         % (label, pi + 1, ci + 1, ech, gsig[:FMT_STRICT], esig[:FMT_STRICT]))
                    break
                if gsig[FMT_STRICT] != esig[FMT_STRICT]:
                    info(scenario, "%s p%d c%d %r: bold %r -> %r"
                         % (label, pi + 1, ci + 1, ech, esig[FMT_STRICT], gsig[FMT_STRICT]))


# ---------------------------------------------------------------- responses

def load_responses(folder):
    found = {}
    for name in sorted(os.listdir(folder)):
        if name.startswith("resp-") and name.endswith(".json"):
            with open(os.path.join(folder, name), "r", encoding="utf-8") as fh:
                raw = fh.read()
            try:
                found[name[5:-5]] = json.loads(raw)
            except Exception as exc:
                fail(name, "response is not JSON: %s (%r)" % (exc, raw[:200]))
    return found


def result_of(scenario, responses, step):
    if step not in responses:
        fail(scenario, "missing response %r (got %s)" % (step, sorted(responses)))
        return {}
    resp = responses[step]
    if not resp.get("ok"):
        fail(scenario, "step %s failed: %s" % (step, json.dumps(resp.get("error"), ensure_ascii=False)))
        return {}
    return resp.get("result", {})


def check_hits(scenario, result, expect, where_filter=None):
    changed = result.get("changed")
    if changed != expect["total"]:
        fail(scenario, "changed=%r, expected %d" % (changed, expect["total"]))
    touched = result.get("shapesTouched")
    if touched != expect["slotsTouched"]:
        fail(scenario, "shapesTouched=%r, expected %d" % (touched, expect["slotsTouched"]))
    hits = result.get("hits") or []
    if len(hits) != expect["total"]:
        fail(scenario, "hits has %d entries, expected %d" % (len(hits), expect["total"]))
    seen = {}
    for hit in hits:
        if hit.get("before") != TOKEN:
            fail(scenario, "hit before=%r, expected %r" % (hit.get("before"), TOKEN))
        key = hit.get("where")
        seen[key] = seen.get(key, 0) + 1
        if where_filter is not None and key != where_filter:
            fail(scenario, "hit landed in scope %r, expected %r" % (key, where_filter))
    if seen != expect["byWhere"]:
        fail(scenario, "hit scopes %r, expected %r" % (seen, expect["byWhere"]))


# ---------------------------------------------------------------- scenarios

DEEP_SCENARIOS = ("full", "dryrun", "restore", "case")


def run_scenario(folder, pristine_slots):
    name = os.path.basename(folder)
    responses = load_responses(folder)
    # deck.plain.pptx is the decrypted copy of what WPS saved; without it there is no
    # file-level evidence to look at (`-SkipDecrypt`).
    plain = os.path.join(folder, "deck.plain.pptx")
    deck_slots = scan(plain) if os.path.exists(plain) else None
    if deck_slots is None and name in DEEP_SCENARIOS:
        info(name, "no deck.plain.pptx: file-level checks skipped (driver ran with -SkipDecrypt)")

    def_e = count_hits(pristine_slots, wheres=DEFAULT_WHERES)   # a bare `replace`
    all_e = count_hits(pristine_slots)                          # every slot, master included

    if name == "dryrun":
        res = result_of(name, responses, "replace")
        check_hits(name, res, def_e)
        if res.get("applied") is not False:
            fail(name, "applied=%r, expected false" % res.get("applied"))
        if res.get("preSnapshot"):
            fail(name, "a dry run must not write a snapshot (%r)" % res.get("preSnapshot"))
        # the driver saves after the dry run, so a write would survive to disk here
        compare_decks(name, pristine_slots, deck_slots)

    elif name == "full":
        res = result_of(name, responses, "replace")
        check_hits(name, res, def_e)
        if res.get("applied") is not True:
            fail(name, "applied=%r, expected true" % res.get("applied"))
        snap = res.get("preSnapshot")
        if not snap or not os.path.exists(snap):
            fail(name, "preSnapshot missing or not on disk: %r" % snap)
        if res.get("skipped"):
            fail(name, "skipped is not empty: %s" % json.dumps(res["skipped"][:3], ensure_ascii=False))
        # the master keeps the token: a bare replace must not touch it
        want = expected_deck(pristine_slots, wheres=DEFAULT_WHERES)
        compare_decks(name, want, deck_slots)

        res2 = result_of(name, responses, "replace-again")
        if res2.get("changed") != 0:
            fail(name, "idempotency: the second run changed=%r, expected 0" % res2.get("changed"))

    elif name == "restore":
        res = result_of(name, responses, "replace")
        check_hits(name, res, def_e)
        back = result_of(name, responses, "restore")
        if back.get("restored", 0) < def_e["slotsTouched"]:
            fail(name, "restore reported %r restored slots, expected >= %d"
                 % (back.get("restored"), def_e["slotsTouched"]))
        compare_decks(name, pristine_slots, deck_slots, strict_format=False)

    elif name == "notes":
        total = all_e["byWhere"].get("notes", 0)
        expect = {"total": total, "byWhere": {"notes": total}, "slotsTouched": 1}
        check_hits(name, result_of(name, responses, "replace"), expect, where_filter="notes")

    elif name == "master":
        total = all_e["byWhere"].get("master", 0)
        expect = {"total": total, "byWhere": {"master": total}, "slotsTouched": 1}
        check_hits(name, result_of(name, responses, "replace"), expect, where_filter="master")

    elif name == "slidesonly":
        total = def_e["byWhere"].get("body", 0)
        expect = {"total": total, "byWhere": {"body": total},
                  "slotsTouched": sum(1 for s in pristine_slots
                                      if s["where"] == "body" and TOKEN in slot_text(s))}
        check_hits(name, result_of(name, responses, "replace"), expect, where_filter="body")

    elif name == "slide1":
        total = def_e["bySlide"].get("1", 0)
        slots = sum(1 for s in pristine_slots
                    if s["slide"] == 1 and s["where"] in DEFAULT_WHERES and TOKEN in slot_text(s))
        expect = {"total": total, "byWhere": {"body": total}, "slotsTouched": slots}
        check_hits(name, result_of(name, responses, "replace"), expect, where_filter="body")

    elif name == "regex":
        total = regex_hits(scoped_slots(pristine_slots), r"20(\d\d)")
        res = result_of(name, responses, "replace")
        if res.get("changed") != total:
            fail(name, "changed=%r, expected %d regex matches" % (res.get("changed"), total))
        for hit in res.get("hits") or []:
            if hit.get("after") != "X" + hit.get("before", "")[2:]:
                fail(name, "regex $1 expansion wrong: %r -> %r" % (hit.get("before"), hit.get("after")))
                break

    elif name == "case":
        res = result_of(name, responses, "replace")
        if res.get("changed") != 0:
            fail(name, "-Case should not match AA2024 when looking for aa2024, changed=%r"
                 % res.get("changed"))
        loose = count_hits(pristine_slots, token="aa2024", wheres=DEFAULT_WHERES)["total"]
        if loose < 1:
            fail(name, "fixture problem: case-insensitive aa2024 should occur at least once")
        compare_decks(name, pristine_slots, deck_slots)

    elif name == "max":
        res = result_of(name, responses, "replace")
        if res.get("changed") != 3:
            fail(name, "-Max 3: changed=%r" % res.get("changed"))
        if not res.get("capped"):
            fail(name, "-Max 3: capped flag missing")

    elif name == "items":
        res = result_of(name, responses, "replace")
        total = def_e["total"] + count_hits(pristine_slots, token="2023", wheres=DEFAULT_WHERES)["total"]
        if res.get("changed") != total:
            fail(name, "items: changed=%r, expected %d (2024 + 2023 rules)" % (res.get("changed"), total))
        if len(res.get("rules") or []) != 2:
            fail(name, "items: rules=%r, expected 2" % res.get("rules"))

    elif name == "footer":
        res = result_of(name, responses, "replace")
        check_hits(name, res, {"total": 1, "byWhere": {"footer": 1}, "slotsTouched": 1},
                   where_filter="footer")
        state_path = os.path.join(folder, "state.json")
        if os.path.exists(state_path):
            with open(state_path, "r", encoding="utf-8") as fh:
                state = json.loads(fh.read())
            if state.get("footerAfter") != "ACME 2025":
                fail(name, "footer after the run is %r, expected 'ACME 2025'"
                     % state.get("footerAfter"))
        else:
            fail(name, "state.json missing (the driver must dump the footer text)")
    else:
        fail(name, "unknown scenario folder")


def main(argv):
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", required=True)
    ap.add_argument("--list", action="store_true")
    args = ap.parse_args(argv)

    root = os.path.abspath(args.dir)
    if not os.path.isdir(root):
        print("no such directory: %s" % root)
        return 2

    folders = []
    for name in sorted(os.listdir(root)):
        folder = os.path.join(root, name)
        if os.path.isdir(folder) and os.path.exists(os.path.join(folder, "pristine.pptx")):
            folders.append(folder)

    for folder in folders:
        slots = scan(os.path.join(folder, "pristine.pptx"))
        print("== %-11s %s" % (os.path.basename(folder),
                               json.dumps(count_hits(slots), ensure_ascii=False)))
    if args.list:
        return 0

    for folder in folders:
        slots = scan(os.path.join(folder, "pristine.pptx"))
        try:
            run_scenario(folder, slots)
        except Exception as exc:            # a broken run must not hide the others
            import traceback
            fail(os.path.basename(folder), "checker crashed: %s" % exc)
            traceback.print_exc()

    for line in _infos:
        print("  info  %s" % line)
    if _failures:
        print("")
        for line in _failures:
            print("  FAIL  %s" % line)
        print("\n%d check(s) failed" % len(_failures))
        return 1
    print("\nall checks passed")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
