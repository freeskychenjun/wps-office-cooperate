#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Assertions for `ppt split` / `ppt merge` (tests/ppt).

The fixture titles every page ("PAGE-1"..."PAGE-12", "A-1"..."C-3"), so a loss, an
extra page or a wrong order shows up in the titles and not only in a count.

  python check_chapters.py --dir <out-chapters>
"""
import argparse
import json
import os
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

_failures = []


def fail(message):
    _failures.append(message)


def load(folder, step):
    path = os.path.join(folder, "resp-%s.json" % step)
    if not os.path.exists(path):
        fail("missing response %r" % step)
        return None, {}
    with open(path, "r", encoding="utf-8") as fh:
        raw = fh.read()
    try:
        doc = json.loads(raw)
    except Exception as exc:
        fail("response %r is not JSON: %s" % (step, exc))
        return None, {}
    return doc, doc.get("result", {})


def titles_of(pptx_path):
    from pptx import Presentation
    prs = Presentation(pptx_path)
    found = []
    for slide in prs.slides:
        text = ""
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                text = shape.text_frame.text.strip()
                break
        found.append(text)
    return found, prs


def master_mark(pptx_path):
    """The marker text in the slide master's title placeholder."""
    from pptx import Presentation
    prs = Presentation(pptx_path)
    for ph in prs.slide_master.placeholders:
        if ph.placeholder_format.type == 1:
            return ph.text_frame.text.strip()
    return ""


def check_split(folder):
    doc, dry = load(folder, "split-dry")
    if doc and doc.get("ok"):
        if not dry.get("dryRun"):
            fail("split -DryRun did not report dryRun")
        files = dry.get("files") or []
        if len(files) != 3:
            fail("split -DryRun planned %d parts, expected 3" % len(files))
        for entry in files:
            if not entry.get("would"):
                fail("split -DryRun entry is not marked would: %r" % entry)

    doc, done = load(folder, "split")
    ok = bool(doc and doc.get("ok"))
    if not ok:
        if doc is not None:
            fail("split failed: %s" % json.dumps(doc.get("error"), ensure_ascii=False))
    else:
        if done.get("chunks") != 3 or done.get("slides") != 12:
            fail("split reported chunks=%r slides=%r, expected 3 / 12"
                 % (done.get("chunks"), done.get("slides")))
        files = done.get("files") or []
        if len(files) != 3:
            fail("split produced %d files, expected 3" % len(files))
        for entry in files:
            if entry.get("slides") != 4:
                fail("part %r has %r slides, expected 4" % (entry.get("chunk"), entry.get("slides")))
            if not os.path.exists(entry.get("out") or ""):
                fail("part file missing: %r" % entry.get("out"))
        expect_first = ["PAGE-1", "PAGE-5", "PAGE-9"]
        for entry, first in zip(files, expect_first):
            if entry.get("first") != first:
                fail("part %r starts at %r, expected %r" % (entry.get("chunk"), entry.get("first"), first))

    # a gap in -Chunks must be refused, not silently dropped
    doc, _ = load(folder, "split-gap")
    if doc is None or doc.get("ok"):
        fail("split with a gap in -Chunks should fail")
    else:
        if "contiguous" not in json.dumps(doc.get("error"), ensure_ascii=False):
            fail("the gap error does not explain the contiguity rule: %s"
                 % json.dumps(doc.get("error"), ensure_ascii=False))

    # the source deck must be untouched by splitting
    doc, read = load(folder, "source-after-split")
    if doc and doc.get("ok"):
        if read.get("to") != 12 and read.get("total") != 12:
            got = read.get("total", read.get("to"))
            if got != 12:
                fail("the source deck no longer has 12 slides after the split (%r)" % got)

    # the parts on disk, read independently
    parts_dir = os.path.join(folder, "parts")
    if not os.path.isdir(parts_dir):
        fail("parts folder missing")
        return
    plains = sorted([n for n in os.listdir(parts_dir) if n.startswith("plain-")])
    if len(plains) != 3:
        fail("expected 3 decrypted part copies, found %d (decrypt step missing?)" % len(plains))
    expect = [["PAGE-%d" % n for n in range(1, 5)],
              ["PAGE-%d" % n for n in range(5, 9)],
              ["PAGE-%d" % n for n in range(9, 13)]]
    # a picture sits on every page whose number is divisible by 3:
    # part 1 -> page 3, part 2 -> page 6, part 3 -> pages 9 and 12
    expect_pics = [1, 1, 2]
    for name, want, pics in zip(plains, expect, expect_pics):
        path = os.path.join(parts_dir, name)
        try:
            got, prs = titles_of(path)
        except Exception as exc:
            fail("could not read %s: %s" % (name, exc))
            continue
        if got != want:
            fail("%s titles are %r, expected %r" % (name, got, want))
        if abs(prs.slide_width.inches - 13.333) > 0.02:
            fail("%s lost the slide width (%.3f in)" % (name, prs.slide_width.inches))
        pictures = sum(1 for slide in prs.slides for s in slide.shapes
                       if s.shape_type is not None and int(s.shape_type) in (13, 11))
        if pictures != pics:
            fail("%s has %d pictures, expected %d" % (name, pictures, pics))
        # the design has to survive the split: pasting pages into a fresh deck throws the
        # master away, copying the file keeps it
        mark = master_mark(path)
        if mark != "MASTER-MARK":
            fail("%s lost the source design: master title reads %r, expected 'MASTER-MARK'"
                 % (name, mark))


def check_merge(folder):
    doc, merged = load(folder, "merge")
    ok = bool(doc and doc.get("ok"))
    if not ok:
        if doc is not None:
            fail("merge failed: %s" % json.dumps(doc.get("error"), ensure_ascii=False))
    else:
        if merged.get("slidesBefore") != 1:
            fail("merge target started with %r slides, expected the 1 title slide"
                 % merged.get("slidesBefore"))
        if merged.get("slidesNow") != 10:
            fail("merge left %r slides, expected 10 (1 host + 3x3)"
                 % merged.get("slidesNow"))
        appended = merged.get("appended") or []
        if len(appended) != 3:
            fail("merge appended %d sources, expected 3" % len(appended))
        for entry in appended:
            if entry.get("slides") != 3:
                fail("merge appended %r slides from %r" % (entry.get("slides"), entry.get("source")))

    doc, read = load(folder, "merge-read")
    if doc and doc.get("ok"):
        text = json.dumps(read, ensure_ascii=False)
        for order in (["A-1", "A-2", "A-3", "B-1"], ["C-1", "C-2", "C-3"]):
            positions = [text.find(marker) for marker in order]
            if -1 in positions:
                fail("merge read is missing %r" % [m for m, p in zip(order, positions) if p < 0])
            elif positions != sorted(positions):
                fail("merged slides are out of order: %r" % order)

    # and the saved deck, read independently
    plain = os.path.join(folder, "host.plain.pptx")
    if not os.path.exists(plain):
        fail("host.plain.pptx missing (decrypt step)")
    else:
        try:
            got, _prs = titles_of(plain)
        except Exception as exc:
            fail("could not read the merged deck: %s" % exc)
            got = []
        want = ["Merged", "A-1", "A-2", "A-3", "B-1", "B-2", "B-3", "C-1", "C-2", "C-3"]
        if got != want:
            fail("merged deck titles are %r, expected %r" % (got, want))

    # -After inserts at a position instead of appending
    doc, dry = load(folder, "merge-after")
    if doc and doc.get("ok"):
        appended = dry.get("appended") or []
        if not appended:
            fail("merge -After -DryRun planned nothing")
        elif appended[0].get("at") != 2:
            fail("merge -After 1 would insert at %r, expected index 2"
                 % appended[0].get("at"))


def main(argv):
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", required=True)
    args = ap.parse_args(argv)
    folder = os.path.abspath(args.dir)
    if not os.path.isdir(folder):
        print("no such directory: %s" % folder)
        return 2

    check_split(folder)
    check_merge(folder)

    if _failures:
        print("")
        for line in _failures:
            print("  FAIL  %s" % line)
        print("\n%d check(s) failed" % len(_failures))
        return 1
    print("\nall checks passed")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
