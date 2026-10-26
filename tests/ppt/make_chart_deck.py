#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Fixture + payloads for `ppt chart` (tests/ppt).

chart-deck.pptx   Scroll 3 pages, 16:9, not a single chart in it.
                  A windowed deck is required: WPS refuses AddChart in a windowless one.
single.json       {"categories":[..],"series":[{"name":..,"values":[..]}]}
twoseries.json    two series (WPS seeds 3, so this has to keep both and drop the third)
line.json         a line chart over 5 points

  python make_chart_deck.py --out <dir>
"""
import argparse
import json
import os
import sys

from pptx import Presentation
from pptx.util import Inches, Pt

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass


def build(path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    titles = ["\u5b9e\u6d4b\u6d41\u91cf\u5bf9\u6bd4",      # 实测流量对比
              "\u4e24\u4e2a\u7cfb\u5217",                  # 两个系列
              "\u6d41\u91cf\u8fc7\u7a0b\u7ebf"]            # 流量过程线
    for text in titles:
        slide = prs.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11), Inches(0.9))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = text
        run.font.size = Pt(32)
        run.font.bold = True
        box.name = "title"
    prs.save(path)


def payloads(out):
    single = {
        "categories": ["2020", "2021", "2022"],
        "series": [{"name": "\u6d41\u91cf", "values": [1200, 1580, 2248]}],
    }
    two = {
        "categories": ["\u4e00\u6708", "\u4e8c\u6708", "\u4e09\u6708"],
        "series": [
            {"name": "\u5b9e\u6d4b", "values": [120, 158, 224]},
            {"name": "\u8bbe\u8ba1", "values": [110, 150, 210]},
        ],
    }
    line = {
        "categories": ["1", "2", "3", "4", "5"],
        "series": [{"name": "\u6c34\u4f4d", "values": [12.1, 12.6, 13.4, 13.0, 12.2]}],
    }
    for name, data in (("single.json", single), ("twoseries.json", two), ("line.json", line)):
        with open(os.path.join(out, name), "w", encoding="utf-8") as fh:
            json.dump(data, fh, ensure_ascii=False, indent=1)


def main(argv):
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True)
    args = ap.parse_args(argv)
    out = os.path.abspath(args.out)
    if not os.path.isdir(out):
        os.makedirs(out)
    build(os.path.join(out, "chart-deck.pptx"))
    payloads(out)
    print(out)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
