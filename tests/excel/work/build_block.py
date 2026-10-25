import openpyxl, json, io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
wb = openpyxl.load_workbook(r'D:\DSH工作区\wps-excel-cooperate\tests\work\waizhou-process.xlsx')
ws = wb.worksheets[0]
rows = list(ws.iter_rows(min_row=1, max_row=3, values_only=True))
for row in rows: print(row)
# build values block: time text, z, q
data = [["时间", "水位(m)", "流量(m3/s)"]]
for r in ws.iter_rows(min_row=2, values_only=True):
    t = r[0]
    ts = t.strftime('%m-%d %H:%M') if hasattr(t, 'strftime') else str(t)
    z = round(float(r[1]), 2) if r[1] is not None else None
    q = float(r[2]) if r[2] is not None else None
    data.append([ts, z, q])
items = {"items": [{"sheet": "Sheet3", "cell": "A7:C142", "values": data}]}
with open(r'D:\DSH工作区\wps-excel-cooperate\tests\work\waizhou-block.json', 'w', encoding='utf-8') as f:
    json.dump(items, f, ensure_ascii=False)
print('rows:', len(data))
